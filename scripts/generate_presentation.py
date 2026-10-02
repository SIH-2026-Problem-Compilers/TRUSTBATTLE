"""Generate the TRUSTBATTLE pitch deck as a .pptx.

Run from repo root:  py scripts/generate_presentation.py
Output:              TRUSTBATTLE_presentation.pptx  (16:9, 7 slides + appendix)

Design: consistent dark theme; every slide uses the same title bar, accent
rule, footer, fonts and colors via the helpers below. Edit the constants at
the top and rerun to tweak.
"""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]

# ---------------------------------------------------------------- theme ----
SLIDE_W, SLIDE_H = 13.333, 7.5              # 16:9
NAVY = RGBColor(0x0B, 0x1D, 0x3A)            # background
PANEL = RGBColor(0x13, 0x2C, 0x52)           # card background
PANEL_DARK = RGBColor(0x0E, 0x22, 0x42)      # title bar
ACCENT = RGBColor(0x00, 0xD1, 0xB2)          # teal
WARN = RGBColor(0xFF, 0xC1, 0x07)            # amber
RED = RGBColor(0xFF, 0x5A, 0x5F)             # alert red
WHITE = RGBColor(0xF5, 0xF7, 0xFA)
GREY = RGBColor(0xA3, 0xB2, 0xC7)
FONT = "Calibri"

TITLE_H = 1.05
MARGIN = 0.55
CONTENT_W = SLIDE_W - 2 * MARGIN
FOOTER_TEXT = ("TRUSTBATTLE — AI-Driven Battlefield Information Integrity & "
               "Trust Assessment Engine")

prs = Presentation()
prs.slide_width = Inches(SLIDE_W)
prs.slide_height = Inches(SLIDE_H)
BLANK = prs.slide_layouts[6]

# -------------------------------------------------------------- helpers ----
def _style_run(run, size, color, bold=False, italic=False):
    f = run.font
    f.name = FONT
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = color


def para(runs, align=PP_ALIGN.LEFT, space_after=4, space_before=0,
         line_spacing=None):
    """runs: list of (text, size, color[, bold[, italic]]) tuples."""
    return {"runs": runs, "align": align, "space_after": space_after,
            "space_before": space_before, "line_spacing": line_spacing}


def add_textbox(slide, x, y, w, h, paras, anchor=MSO_ANCHOR.TOP, wrap=True):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    for i, p in enumerate(paras):
        paragraph = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        paragraph.alignment = p["align"]
        paragraph.space_after = Pt(p["space_after"])
        paragraph.space_before = Pt(p["space_before"])
        if p["line_spacing"]:
            paragraph.line_spacing = p["line_spacing"]
        for run_spec in p["runs"]:
            text, size, color = run_spec[0], run_spec[1], run_spec[2]
            bold = run_spec[3] if len(run_spec) > 3 else False
            italic = run_spec[4] if len(run_spec) > 4 else False
            run = paragraph.add_run()
            run.text = text
            _style_run(run, size, color, bold, italic)
    return box


def add_card(slide, x, y, w, h, fill=PANEL, line=None, radius=0.07):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x),
                                  Inches(y), Inches(w), Inches(h))
    card.adjustments[0] = radius
    card.fill.solid()
    card.fill.fore_color.rgb = fill
    if line is not None:
        card.line.color.rgb = line
        card.line.width = Pt(1.0)
    else:
        card.line.fill.background()
    card.shadow.inherit = False
    return card


def add_slide(title, kicker):
    slide = prs.slides.add_slide(BLANK)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NAVY

    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(SLIDE_W),
                                 Inches(TITLE_H))
    bar.fill.solid()
    bar.fill.fore_color.rgb = PANEL_DARK
    bar.line.fill.background()
    bar.shadow.inherit = False

    rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0,
                                  Inches(TITLE_H - 0.04), Inches(SLIDE_W),
                                  Inches(0.05))
    rule.fill.solid()
    rule.fill.fore_color.rgb = ACCENT
    rule.line.fill.background()
    rule.shadow.inherit = False

    add_textbox(slide, MARGIN, 0.10, CONTENT_W - 2.6, TITLE_H - 0.22,
                [para([(title, 26, WHITE, True)])], anchor=MSO_ANCHOR.MIDDLE)
    add_textbox(slide, SLIDE_W - 2.9, 0.10, 2.35, TITLE_H - 0.22,
                [para([(kicker, 11, ACCENT, True)], align=PP_ALIGN.RIGHT)],
                anchor=MSO_ANCHOR.MIDDLE)
    add_textbox(slide, MARGIN, SLIDE_H - 0.36, CONTENT_W, 0.28,
                [para([(FOOTER_TEXT, 9, GREY, False)])])
    return slide


def label(slide, x, y, w, text, color=ACCENT, size=11):
    return add_textbox(slide, x, y, w, 0.28,
                       [para([(text, size, color, True)])])


def bullets(slide, x, y, w, h, items, size=11, gap=5, color=WHITE,
            line_spacing=1.05):
    paras = []
    for it in items:
        if isinstance(it, tuple):
            runs = [("•  ", size, ACCENT, True), (it[0], size, color, True),
                    (it[1], size, color, False)]
        else:
            runs = [("•  ", size, ACCENT, True), (it, size, color, False)]
        paras.append(para(runs, space_after=gap, line_spacing=line_spacing))
    return add_textbox(slide, x, y, w, h, paras)


def flow_box(slide, x, y, w, h, lines, fill=PANEL, line=ACCENT, radius=0.14):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x),
                                 Inches(y), Inches(w), Inches(h))
    shp.adjustments[0] = radius
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = line
    shp.line.width = Pt(1.0)
    shp.shadow.inherit = False
    tf = shp.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = tf.margin_right = Inches(0.06)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    for i, (text, size, color, bold) in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = text
        _style_run(r, size, color, bold)
    return shp


def down_arrow(slide, cx, y, h=0.14, color=GREY):
    shp = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(cx - 0.07),
                                 Inches(y), Inches(0.14), Inches(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


# ------------------------------------------------------- slide 1 · track ----
s = add_slide("Problem Statement Track", "01 / 07")
rows = [
    ("PROBLEM STATEMENT TITLE",
     "[ Paste the exact PS title + PS ID from the official portal — idea maps "
     "to themes on GNSS spoofing / assured sensor-information integrity ]"),
    ("THEME", "Aerospace & Defence  ·  confirm against the official PS"),
    ("PS CATEGORY", "Software"),
    ("TEAM ID", "[ your Team ID ]"),
    ("TEAM NAME", "[ your Team Name ]"),
]
y = 1.42
for lab, val in rows:
    add_card(s, MARGIN, y, CONTENT_W, 0.86)
    add_textbox(s, MARGIN + 0.28, y, 2.9, 0.86,
                [para([(lab, 11, ACCENT, True)])], anchor=MSO_ANCHOR.MIDDLE)
    add_textbox(s, MARGIN + 3.3, y + 0.06, CONTENT_W - 3.6, 0.74,
                [para([(val, 13.5, WHITE, False)])], anchor=MSO_ANCHOR.MIDDLE)
    y += 1.0
add_textbox(s, MARGIN, y + 0.02, CONTENT_W, 0.3,
            [para([("Fill Team ID, Team Name and the official PS title / ID "
                    "exactly as on the portal before exporting.", 10, GREY,
                    False, True)])])

# --------------------------------------------------------- slide 2 · idea ----
s = add_slide("Idea", "02 / 07")
add_card(s, MARGIN, 1.38, CONTENT_W, 1.06)
label(s, MARGIN + 0.28, 1.50, 6.0, "PROJECT / IDEA NAME")
add_textbox(s, MARGIN + 0.28, 1.76, CONTENT_W - 0.6, 0.62, [
    para([("TRUSTBATTLE", 24, WHITE, True),
          ("   —  AI-Driven Battlefield Information Integrity & Trust "
           "Assessment Engine", 13, GREY, False)])])

add_textbox(s, MARGIN, 2.60, CONTENT_W, 0.55,
            [para([("“Don't just ask what the battlefield sensors are "
                    "reporting — ask whether the information they report can "
                    "be trusted.”", 14.5, ACCENT, False, True)],
                  align=PP_ALIGN.CENTER)], anchor=MSO_ANCHOR.MIDDLE)

add_card(s, MARGIN, 3.30, 5.85, 1.62)
label(s, MARGIN + 0.26, 3.42, 5.3, "PROBLEM BEING ADDRESSED")
add_textbox(s, MARGIN + 0.26, 3.70, 5.35, 1.12, [
    para([("UAV, GNSS/IMU, visual and network data can be spoofed, replayed, "
           "injected, jammed or simply faulty — and a manipulated measurement "
           "can look perfectly legitimate. Operators see ", 11.5, WHITE),
          ("what", 11.5, WHITE, True, True),
          (" sensors report, but not whether each individual observation is "
           "trustworthy right now.", 11.5, WHITE)], space_after=0,
         line_spacing=1.08)])

add_card(s, 6.62, 3.30, 6.16, 1.62)
label(s, 6.88, 3.42, 5.6, "PROPOSED SOLUTION")
add_textbox(s, 6.88, 3.70, 5.64, 1.12, [
    para([("An AI-assisted ", 11.5, WHITE),
          ("Information Integrity Layer", 11.5, WHITE, True),
          (" between sensors and fusion. Every observation is scored 0–100 "
           "from five independent evidence streams — physical, temporal, "
           "cyber/network, cross-sensor and historical — suspicious sources "
           "are down-weighted in fusion, and the operator sees a "
           "plain-language evidence checklist.", 11.5, WHITE)], space_after=0,
         line_spacing=1.08)])

label(s, MARGIN, 5.14, 6.0, "KEY DIFFERENTIATORS")
chips = [
    ("Explainable evidence",
     "Checklist of passing / failing checks — not a black-box score"),
    ("Dynamic trust",
     "Drops during anomalies, recovers after — no permanent blacklisting"),
    ("Trust-aware fusion",
     "Suspicious sources automatically down-weighted in the estimate"),
    ("Honest scope",
     "Integrity assessment — never claims “hacked” without sufficient proof"),
]
cx = MARGIN
for title_, sub in chips:
    add_card(s, cx, 5.46, 2.92, 1.28, fill=PANEL, line=ACCENT)
    add_textbox(s, cx + 0.18, 5.58, 2.56, 1.06, [
        para([(title_, 12, ACCENT, True)], space_after=3),
        para([(sub, 9.5, GREY, False)], space_after=0, line_spacing=1.05)])
    cx += 3.10

# ------------------------------------------- slide 3 · technical approach ----
s = add_slide("Technical Approach", "03 / 07")

label(s, MARGIN, 1.34, 5.0, "TECHNOLOGIES / TOOLS USED")
blocks = [
    ("FRONTEND & BACKEND", [
        "React + Vite · Leaflet/Mapbox · Recharts/Plotly (dashboard)",
        "Python FastAPI · REST + WebSocket real-time streaming",
        "PostgreSQL (SQLite in dev) via SQLAlchemy"]),
    ("AI / ML LIBRARIES & MODELS", [
        "scikit-learn Isolation Forest — primary anomaly detector",
        "One-Class SVM baseline · IF+OCSVM ensemble candidate",
        "pandas / NumPy pipelines · LSTM / autoencoder on roadmap"]),
    ("CLOUD & DEPLOYMENT TOOLS", [
        "Modular Python packages driven by shared configs/settings.yaml",
        "Deployable on any cloud VM / container platform (e.g. Docker)"]),
    ("VERSION CONTROL & CI/CD", [
        "Git / GitHub · per-module pytest suites (15 test files)",
        "Frozen data contracts + change-request workflow (docs/contracts)"]),
]
by = 1.70
for head, items in blocks:
    label(s, MARGIN, by, 5.05, head, color=GREY, size=10.5)
    bullets(s, MARGIN, by + 0.26, 5.05, 0.95, items, size=10.5)
    by += 1.24

DX, DW = 5.95, 6.83
DCX = DX + DW / 2
label(s, DX, 1.34, DW, "SYSTEM ARCHITECTURE")
flow_box(s, DX + 0.7, 1.70, DW - 1.4, 0.55,
         [("BATTLEFIELD SOURCES — GNSS · IMU · Visual · Telemetry/Network",
           11, WHITE, True)])
down_arrow(s, DCX, 2.29)
flow_box(s, DX + 1.4, 2.47, DW - 2.8, 0.40,
         [("FEATURE EXTRACTION", 11, WHITE, True)])
down_arrow(s, DCX, 2.91)
aw = 2.18
for i, (name, sub) in enumerate([
        ("M1 · PHYSICAL", "state consistency"),
        ("M2 · TEMPORAL", "replay / timing checks"),
        ("M4 · CYBER", "network anomaly scenarios")]):
    ax = DX + i * (aw + 0.145)
    flow_box(s, ax, 3.09, aw, 0.75,
             [(name, 10.5, WHITE, True), (sub, 9, GREY, False)],
             line=WARN if i == 2 else ACCENT)
down_arrow(s, DCX, 3.88)
flow_box(s, DX + 1.4, 4.06, DW - 2.8, 0.40,
         [("EVIDENCE ENGINE", 11, WHITE, True)])
down_arrow(s, DCX, 4.50)
flow_box(s, DX + 0.9, 4.68, DW - 1.8, 0.50,
         [("OBSERVATION TRUST MODEL (0–100)", 11.5, ACCENT, True)])
down_arrow(s, DCX, 5.22)
ow = 3.30
flow_box(s, DX, 5.40, ow, 0.62,
         [("TRUSTED", 11, WHITE, True), ("→ robust fusion", 9.5, GREY, False)],
         line=ACCENT)
flow_box(s, DX + DW - ow, 5.40, ow, 0.62,
         [("SUSPICIOUS", 11, WHITE, True),
          ("→ verification alert", 9.5, GREY, False)], line=RED)
down_arrow(s, DCX, 6.06)
flow_box(s, DX + 1.4, 6.24, DW - 2.8, 0.42,
         [("EXPLAINABLE DASHBOARD (M5)", 11, WHITE, True)])

# --------------------------------------- slide 4 · feasibility & viability ----
s = add_slide("Feasibility & Viability", "04 / 07")
colw, gap = 3.96, 0.175
cols = [
    ("TECHNICAL FEASIBILITY", [
        "Fully open-source, mature stack — pipeline already runs end-to-end "
        "(integration/run_demo.py)",
        "No classified data: public GNSS spoofing datasets (TEXBAT) + "
        "synthetic attacks with known ground truth",
        "Unsupervised Isolation Forest — no large labelled dataset needed"]),
    ("OPERATIONAL FEASIBILITY", [
        "Adds one layer on top of existing sensor fusion — no sensor "
        "replacement, no workflow disruption",
        "Operator-friendly output: trust %, green/amber/red states, "
        "plain-English evidence checklist",
        "Alert thresholds are calibrated, tunable and explainable"]),
    ("ECONOMIC FEASIBILITY", [
        "Zero licence cost — 100% open-source stack",
        "Lightweight tree models train in seconds — commodity hardware, no "
        "GPU farm",
        "Simulation-first evaluation avoids costly physical test ranges"]),
]
cx = MARGIN
for head, items in cols:
    add_card(s, cx, 1.42, colw, 2.55)
    label(s, cx + 0.22, 1.56, colw - 0.4, head)
    bullets(s, cx + 0.22, 1.88, colw - 0.44, 2.0, items, size=10.5, gap=6)
    cx += colw + gap

add_card(s, MARGIN, 4.16, CONTENT_W, 1.62)
label(s, MARGIN + 0.24, 4.30, 8.0, "SCALABILITY & SUSTAINABILITY")
bullets(s, MARGIN + 0.24, 4.62, CONTENT_W - 0.5, 1.1, [
    ("Config-driven (YAML): ", "thresholds, weights & check limits retunable "
     "without code changes"),
    ("Maintainable: ", "strictly separated modules, frozen data contracts + "
     "change-request governance"),
    ("Expansion roadmap: ", "V2 +Camera/EO → V3 +Radar → V4 multi-UAV → "
     "V5 edge deployment & real hardware"),
], size=11, gap=6)

add_card(s, MARGIN, 5.98, CONTENT_W, 0.92, fill=PANEL, line=WARN)
add_textbox(s, MARGIN + 0.24, 6.08, CONTENT_W - 0.5, 0.74, [
    para([("PROOF POINT (controlled simulation, M2 module):  ", 11, WARN,
           True),
          ("combined detector F1 = 1.00 across replay / "
           "telemetry-manipulation / network-anomaly scenarios, 0 false "
           "positives on 228 clean windows, detection latency ≤ 5 s "
           "(synthetic data — labelled as such).", 11, WHITE, False)],
         space_after=0, line_spacing=1.1)])

# ---------------------------------------------------- slide 5 · impact ----
s = add_slide("Impact & Benefits", "05 / 07")
cards = [
    ("1", "USER BENEFITS", [
        "Instant answer: “can I trust this observation now?”",
        "0–100 score + evidence + recommended action"]),
    ("2", "SOCIAL IMPACT", [
        "Crew & mission safety — decisions rest on verified information",
        "Relevant to GNSS spoofing seen in real conflicts"]),
    ("3", "ECONOMIC IMPACT", [
        "Software-only layer over existing sensors",
        "Prevents costly corrupted-data decisions; near-zero running cost"]),
    ("4", "ENVIRONMENTAL IMPACT", [
        "Simulation-first development — no physical test ranges",
        "Efficient models suit low-power edge hardware"]),
    ("5", "EFFICIENCY & PRODUCTIVITY", [
        "Automates sensor-veracity cross-checking",
        "≤5 s detection limits corruption spread; auto trust recovery"]),
]
cw, cgap = 2.35, 0.12
cx = MARGIN
for num, head, items in cards:
    add_card(s, cx, 1.42, cw, 2.42)
    badge = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx + 0.16),
                               Inches(1.56), Inches(0.4), Inches(0.4))
    badge.fill.solid()
    badge.fill.fore_color.rgb = ACCENT
    badge.line.fill.background()
    badge.shadow.inherit = False
    btf = badge.text_frame
    btf.margin_left = btf.margin_right = btf.margin_top = btf.margin_bottom = 0
    br = btf.paragraphs[0].add_run()
    br.text = num
    _style_run(br, 14, NAVY, True)
    add_textbox(s, cx + 0.66, 1.56, cw - 0.78, 0.44,
                [para([(head, 10.5, WHITE, True)])],
                anchor=MSO_ANCHOR.MIDDLE)
    bullets(s, cx + 0.18, 2.12, cw - 0.36, 1.62, items, size=9.5, gap=5)
    cx += cw + cgap

label(s, MARGIN, 4.12, 6.0, "IMPACT PATHWAY")
steps = [
    ("PROBLEM", "Unverifiable sensor data under spoofing / jamming / "
                "malfunction / deception"),
    ("SOLUTION", "Multi-evidence integrity layer + dynamic trust + "
                 "trust-aware fusion"),
    ("IMMEDIATE IMPACT", "Suspicious source down-weighted; evidence-backed "
                         "alert; fusion stays accurate"),
    ("LONG-TERM IMPACT", "Resilient, explainable information systems — "
                         "transferable to AVs, drones, critical infrastructure"),
]
sx, sw, step = MARGIN, 3.25, 3.0
for i, (head, sub) in enumerate(steps):
    ch = s.shapes.add_shape(MSO_SHAPE.CHEVRON, Inches(sx), Inches(4.46),
                            Inches(sw), Inches(0.78))
    ch.fill.solid()
    ch.fill.fore_color.rgb = PANEL
    ch.line.color.rgb = ACCENT if i < 3 else WARN
    ch.line.width = Pt(1.0)
    ch.shadow.inherit = False
    ctf = ch.text_frame
    ctf.margin_left = Inches(0.25)
    ctf.margin_right = Inches(0.1)
    ctf.margin_top = ctf.margin_bottom = 0
    cr = ctf.paragraphs[0].add_run()
    cr.text = head
    _style_run(cr, 12, ACCENT if i < 3 else WARN, True)
    ctf.paragraphs[0].alignment = PP_ALIGN.CENTER
    add_textbox(s, sx + 0.15, 5.36, sw - 0.25, 1.0,
                [para([(sub, 9.5, GREY, False)], align=PP_ALIGN.CENTER,
                      space_after=0, line_spacing=1.05)])
    sx += step

# ------------------------------------------- slide 6 · research & references ----
s = add_slide("Research & References", "06 / 07")
label(s, MARGIN, 1.36, 6.0, "RESEARCH / LITERATURE SURVEY")
groups = [
    ("Programs & doctrine", [
        "U.S. Army SBIR program: “Ensuring Sensor Data Security and "
        "Integrity” — protecting sensor data across its lifecycle",
        "India's SANJAY Battlefield Surveillance System — integrates "
        "ground/aerial sensors and confirms veracity before the common "
        "picture"]),
    ("Existing solutions & products", [
        "Anti-jam GNSS receivers / GPS firewalls, network intrusion "
        "detection, Kalman-style fusion — each covers one fragment",
        "We combine physical + temporal + cyber + cross-sensor evidence "
        "into one observation-level trust score with explainability"]),
    ("Documented threats", [
        "GNSS spoofing & jamming against UAVs in the Russia–Ukraine "
        "conflict",
        "Decoys / electronic deception produce “valid but wrong” sensor "
        "observations"]),
    ("Datasets used for training / testing", [
        "UT Austin TEXBAT — public GPS spoofing test battery",
        "Our synthetic UAV telemetry with ground-truth-labelled attacks "
        "(spoof / replay / telemetry-manip / network anomaly) at 0–30% "
        "corruption levels"]),
]
gy = 1.68
for head, items in groups:
    label(s, MARGIN, gy, 6.05, head, color=WHITE, size=11.5)
    bullets(s, MARGIN, gy + 0.26, 6.05, 0.95, items, size=10, gap=4)
    gy += 1.24

RX, RW = 6.95, 5.83
label(s, RX, 1.36, RW, "REFERENCES")
refs = [
    ("Liu, Ting & Zhou — “Isolation Forest”, ", "IEEE ICDM, 2008."),
    ("Schölkopf et al. — “Estimating the Support of a High-Dimensional "
     "Distribution”, ", "Neural Computation, 2001 (One-Class SVM)."),
    ("Humphreys et al. — UT Austin Radionavigation Lab: GPS civil-spoofer "
     "studies & TEXBAT dataset, ", "radionavlab.ae.utexas.edu."),
    ("U.S. Army SBIR topic: “Ensuring Sensor Data Security and Integrity”, ",
     "armysbir.com."),
    ("Press Information Bureau — official releases on India's SANJAY "
     "Battlefield Surveillance System, ", "pib.gov.in."),
    ("Official documentation: ", "scikit-learn, FastAPI, React/Vite, "
     "Leaflet/Mapbox, Plotly/Recharts, WebSocket API."),
]
rparas = []
for i, (a, b) in enumerate(refs, start=1):
    rparas.append(para([(f"{i}.  ", 11, ACCENT, True), (a, 11, WHITE, False),
                        (b, 11, GREY, False)], space_after=9,
                       line_spacing=1.1))
add_textbox(s, RX, 1.72, RW, 4.9, rparas)

# ------------------------------------------ slide 7 · important instructions ----
s = add_slide("Important Instructions", "07 / 07")
items = [
    ("Use the real numbers, labelled honestly — ",
     "quote F1 / FPR / latency as “controlled simulation results”, never as "
     "field-tested claims."),
    ("Honesty guardrails — ",
     "never claim “sensors are being hacked”, “nobody has solved this”, or "
     "“our AI proves the attack”."),
    ("Diagrams over prose — ",
     "layered architecture, trust drop/recovery curve (94% → 31% → 94%), M2 "
     "timeline graphs (appendix)."),
    ("One-line demo story — ",
     "Normal (green, 94%) → spoofing injected → trust 31% + evidence → GNSS "
     "down-weighted → recovery to 94%."),
    ("Consistent formatting — ",
     "one font family, 2–3 colours, max ~6 bullets per slide, no decorative "
     "clip-art."),
    ("Accuracy — ",
     "Team ID, PS title/ID and team names must match the official portal "
     "before export."),
    ("Concise & relevant — ",
     "every bullet earns its place; cut anything that doesn't answer what / "
     "why / how / so what."),
]
iy = 1.42
for head, rest in items:
    add_card(s, MARGIN, iy, CONTENT_W, 0.70)
    add_textbox(s, MARGIN + 0.24, iy, 0.45, 0.70,
                [para([("✓", 15, ACCENT, True)])], anchor=MSO_ANCHOR.MIDDLE)
    add_textbox(s, MARGIN + 0.72, iy, CONTENT_W - 1.0, 0.70,
                [para([(head, 12, WHITE, True), (rest, 12, GREY, False)],
                      space_after=0, line_spacing=1.05)],
                anchor=MSO_ANCHOR.MIDDLE)
    iy += 0.78

# ------------------------------------------------ appendix · M2 graphs ----
s = add_slide("Appendix — M2 Evaluation Graphs", "APPENDIX")
graph_dir = ROOT / "docs" / "reports" / "m2_graphs"
graphs = [
    ("m2_score_distributions.png", "Anomaly score distributions"),
    ("m2_roc_pr.png", "ROC & precision-recall curves"),
    ("m2_timeline_fallback_replay.png", "Replay attack detection timeline"),
    ("m2_timeline_fallback_network_anomaly.png",
     "Network-anomaly detection timeline"),
]
cell_w, cell_h = 5.9, 2.62
positions = [(0.65, 1.45), (6.75, 1.45), (0.65, 4.35), (6.75, 4.35)]
for (fname, caption), (gx, gy2) in zip(graphs, positions):
    path = graph_dir / fname
    if path.exists():
        pic = s.shapes.add_picture(str(path), Inches(gx), Inches(gy2),
                                   height=Inches(cell_h - 0.35))
        pw = pic.width / 914400
        pic.left = Inches(gx + (cell_w - pw) / 2)
    else:
        add_card(s, gx, gy2, cell_w, cell_h - 0.35, fill=PANEL_DARK)
        add_textbox(s, gx, gy2 + 0.8, cell_w, 0.5,
                    [para([(f"[ missing: {fname} ]", 10, GREY, False)],
                          align=PP_ALIGN.CENTER)])
    add_textbox(s, gx, gy2 + cell_h - 0.30, cell_w, 0.28,
                [para([(caption, 9.5, GREY, False)], align=PP_ALIGN.CENTER)])
add_textbox(s, MARGIN, 7.06, CONTENT_W, 0.3, [
    para([("Controlled-simulation results — regenerate with "
           "py -m member2_temporal.src.evaluate when real M4 scenario data "
           "lands.", 9.5, GREY, False, True)])])

# ----------------------------------------------------------------- save ----
OUT = ROOT / "TRUSTBATTLE_presentation.pptx"
prs.save(OUT)
print(f"saved {OUT.name}  ({OUT.stat().st_size / 1024:.0f} KB, "
      f"{len(prs.slides)} slides)")
